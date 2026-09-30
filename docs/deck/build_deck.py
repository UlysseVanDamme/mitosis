"""Build the Mitosis pitch deck (5 slides) as mitosis-pitch.pptx.

Run from anywhere:  docs/deck/.venv/bin/python docs/deck/build_deck.py
Needs python-pptx and Pillow (install into docs/deck/.venv with uv).
"""
from pathlib import Path
import tempfile

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
SHOTS = HERE.parent / "screenshots"
OUT = HERE / "mitosis-pitch.pptx"
TMP = Path(tempfile.mkdtemp(prefix="mitosis-deck-"))

# Palette taken from the app's dark microscope UI.
BG = "071424"
PANEL = "0D1E33"
PANEL2 = "12263F"
LINE = "243B57"
TEXT = "E8EEF5"
MUTED = "8A9AB0"
CYAN = "5EE6D0"
GREEN = "4ADE80"
RED = "F2555A"
AMBER = "F5B84B"
VIOLET = "B18CF0"
FONT = "Arial"

prs = Presentation()
prs.slide_width = Emu(12192000)
prs.slide_height = Emu(6858000)
BLANK = prs.slide_layouts[6]


def rgb(h):
    return RGBColor.from_string(h)


def new_slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(BG)
    return s


def text(slide, x, y, w, h, runs, size=18, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None):
    """runs: str, or list of paragraphs; each paragraph is str or list of (text, overrides)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        parts = para if isinstance(para, list) else [(para, {})]
        for t, o in parts:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = rgb(o.get("color", color))
    return tb


def shape(slide, kind, x, y, w, h, fill=None, line=None, lw=1.0):
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def card(slide, x, y, w, h, fill=PANEL, line=LINE):
    sh = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line, 1)
    sh.adjustments[0] = 0.06
    return sh


def cell(slide, cx, cy, r, color, fill=PANEL2, lw=2.5, label=None, lsize=14):
    """A 'cell': circle with a coloured membrane, the app's motif."""
    c = shape(slide, MSO_SHAPE.OVAL, cx - r, cy - r, 2 * r, 2 * r, fill, color, lw)
    if label:
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = label
        run.font.name = FONT
        run.font.size = Pt(lsize)
        run.font.bold = True
        run.font.color.rgb = rgb(color)
    return c


def connector(slide, x1, y1, x2, y2, color=LINE, lw=1.5, arrow=False):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(lw)
    if arrow:
        # python-pptx has no arrowhead API: add a tailEnd to the line XML.
        from pptx.oxml.ns import qn
        ln_el = ln.line._get_or_add_ln()
        tail = ln_el.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
        ln_el.append(tail)
    return ln


def kicker(slide, n, label):
    text(slide, 0.7, 0.55, 8, 0.3, [[(f"0{n}", {"color": CYAN, "bold": True}),
                                    ("   " + label.upper(), {"color": MUTED})]], size=12)


def crop(src, box, name):
    out = TMP / name
    Image.open(src).crop(box).save(out, quality=92)
    return out


# ---------------------------------------------------------------- slide 1
s = new_slide()
text(s, 0.7, 0.55, 8, 0.3, [[("MITOSIS", {"color": CYAN, "bold": True}),
                            ("   Monday, 8:40. Sofie inherits Brouwerij Van Dessel.", {"color": MUTED})]], size=12)
text(s, 0.7, 1.35, 12, 1.2, "Three documents. Three answers.", size=46, bold=True)
text(s, 0.7, 2.45, 11, 0.5, "“Which indexation applies in January?”", size=22, color=MUTED)

docs = [
    ("2.13%", "Forecast", "Pro-Pay, 14 Oct", AMBER),
    ("2.21%", "Final figure", "Agoria, 22 Dec", GREEN),
    ("February", "Company agreement", "Van Dessel CAO + Teams", RED),
]
x0, w, gap, y = 0.7, 3.75, 0.45, 3.55
for i, (val, kind, src, col) in enumerate(docs):
    x = x0 + i * (w + gap)
    card(s, x, y, w, 2.6)
    # folded-corner document glyph
    shape(s, MSO_SHAPE.FOLDED_CORNER, x + 0.4, y + 0.4, 0.42, 0.52, PANEL2, col, 1.5)
    text(s, x + 1.0, y + 0.42, w - 1.2, 0.5, kind, size=14, color=MUTED)
    text(s, x + 1.0, y + 0.66, w - 1.2, 0.4, src, size=12, color=MUTED)
    text(s, x + 0.4, y + 1.25, w - 0.8, 1.0, val, size=48, bold=True, color=col)
text(s, 0.7, 6.55, 12, 0.4, "Her search finds all three. Nothing tells her which one to trust.",
     size=16, color=MUTED)
s.notes_slide.notes_text_frame.text = (
    "Sofie just inherited the Brouwerij Van Dessel portfolio, and the first client question is already "
    "waiting: what indexation in January? Her search finds three documents and a Teams message. "
    "They say 2.13%, 2.21% and 'February'.")

# ---------------------------------------------------------------- slide 2
s = new_slide()
kicker(s, 2, "Three ways to hold knowledge")
text(s, 0.7, 0.95, 12, 0.9, "Dead, connected, alive.", size=44, bold=True)

compare = SHOTS / "compare-act4.png"
if compare.exists():
    s.shapes.add_picture(str(compare), Inches(0.7), Inches(2.0), width=Inches(11.93))
else:
    cols = [
        ("Dead", "Classic AI search", "Blends chunks into one confident answer. A forecast looks as relevant as the final figure.", RED),
        ("Connected", "Knowledge graph", "Someone has to draw the map and keep it current. It goes stale the week after.", AMBER),
        ("Alive", "Mitosis", "Divides itself as it reads, and catches the contradiction before anyone asks.", GREEN),
    ]
    w, gap, y, h = 3.75, 0.45, 2.1, 4.7
    for i, (name, what, fault, col) in enumerate(cols):
        x = 0.7 + i * (w + gap)
        card(s, x, y, w, h, fill=PANEL if i < 2 else "0B2A2E", line=LINE if i < 2 else CYAN)
        # illustration zone
        ix, iy, iw, ih = x + 0.3, y + 0.3, w - 0.6, 2.1
        if i == 0:  # scattered chunks, no structure
            import random
            random.seed(4)
            for k in range(14):
                cw = random.uniform(0.35, 0.8)
                shape(s, MSO_SHAPE.RECTANGLE, ix + random.uniform(0, iw - cw), iy + random.uniform(0.1, ih - 0.3),
                      cw, 0.16, LINE, None)
        elif i == 1:  # fixed graph
            pts = [(0.3, 0.5), (1.2, 0.25), (2.2, 0.6), (2.9, 1.3), (1.7, 1.4), (0.6, 1.6), (2.4, 1.9)]
            edges = [(0, 1), (1, 2), (2, 3), (1, 4), (4, 3), (0, 5), (5, 4), (4, 6), (3, 6)]
            for a, b in edges:
                connector(s, ix + pts[a][0], iy + pts[a][1], ix + pts[b][0], iy + pts[b][1], MUTED, 1.25)
            for px, py in pts:
                shape(s, MSO_SHAPE.OVAL, ix + px - 0.1, iy + py - 0.1, 0.2, 0.2, PANEL2, MUTED, 1.5)
        else:  # real swarm crop
            img = crop(SHOTS / "1920-03-swarm.jpg", (380, 150, 1060, 860), "swarm.jpg")
            s.shapes.add_picture(str(img), Inches(ix + (iw - 2.1) / 2), Inches(iy), height=Inches(2.1))
        text(s, x + 0.3, y + 2.65, w - 0.6, 0.6, name, size=30, bold=True, color=col)
        text(s, x + 0.3, y + 3.2, w - 0.6, 0.4, what, size=14, color=MUTED)
        text(s, x + 0.3, y + 3.6, w - 0.6, 1.0, fault, size=14, color=TEXT)
s.notes_slide.notes_text_frame.text = (
    "Classic AI search retrieves a handful of chunks and blends them into one confident answer: a forecast "
    "looks exactly as relevant as the final figure. A knowledge graph connects things, but somebody has to "
    "draw that map and keep it current. So we built something that reads everything first. It starts as one "
    "agent, and when it knows too much to keep in mind at once, it divides like a cell.")

# ---------------------------------------------------------------- slide 3
s = new_slide()
kicker(s, 3, "How it works")
text(s, 0.7, 0.95, 12, 0.9, "It reads first, so you don’t have to search.", size=40, bold=True)

steps = [
    ("Read", "every document, chat and ticket", CYAN, "1"),
    ("Divide", "by joint committee, client, country", VIOLET, "2"),
    ("Catch", "specialists see their whole domain", RED, "3"),
    ("Tell the owner", "Jan settles it in one click", AMBER, "4"),
    ("Verified", "everyone gets the green answer", GREEN, "5"),
]
n = len(steps)
cy, r = 3.55, 0.72
xs = [0.7 + r + i * ((13.333 - 1.4 - 2 * r) / (n - 1)) for i in range(n)]
for i in range(n - 1):
    connector(s, xs[i] + r + 0.12, cy, xs[i + 1] - r - 0.12, cy, MUTED, 1.75, arrow=True)
for i, (name, sub, col, num) in enumerate(steps):
    cx = xs[i]
    if i == 1:  # dividing cell: two lobes
        cell(s, cx - 0.22, cy, r * 0.78, col, lw=2.5)
        cell(s, cx + 0.22, cy, r * 0.78, col, lw=2.5)
        text(s, cx - 0.5, cy - 0.25, 1.0, 0.5, num, size=24, bold=True, color=col,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    else:
        fill = "3A1418" if i == 2 else ("0F3322" if i == 4 else PANEL2)
        cell(s, cx, cy, r, col, fill=fill, lw=3 if i in (2, 4) else 2.5, label=num, lsize=24)
    text(s, cx - 1.2, cy + r + 0.35, 2.4, 0.45, name, size=20, bold=True, align=PP_ALIGN.CENTER)
    text(s, cx - 1.2, cy + r + 0.85, 2.4, 0.7, sub, size=14, color=MUTED, align=PP_ALIGN.CENTER)
text(s, 0.7, 6.55, 12, 0.4,
     [[("85%", {"color": CYAN, "bold": True}),
       (" of routing is rules and vector math, no LLM call. Facts are extracted once, when a document arrives.",
        {"color": MUTED})]], size=14)
s.notes_slide.notes_text_frame.text = (
    "It splits by joint committee, by client and by country, and it writes down why. Nobody drew this map "
    "by hand. Each specialist keeps its whole domain in view, so it notices when two sources disagree while "
    "it is still reading. Here it's a forecast of 2.13% against the final 2.21%. The swarm sends Jan, who "
    "owns PC 200, only the questions he can settle, and shows what they affect: one client config still "
    "uses the old number. He confirms it, and from then on everyone gets the verified answer.")

# ---------------------------------------------------------------- slide 4
s = new_slide()
kicker(s, 4, "Does it work")
text(s, 0.7, 0.95, 12, 0.9, "Tested on 102 documents with planted traps.", size=40, bold=True)

# hero: accuracy comparison
card(s, 0.7, 2.1, 5.6, 4.1)
text(s, 1.1, 2.4, 4.8, 0.4, "Answer accuracy, 18 golden questions", size=14, color=MUTED)
text(s, 1.1, 2.85, 4.8, 1.3, [[("78%", {"color": GREEN}), ("  vs 61%", {"color": MUTED, "size": 30})]],
     size=72, bold=True)
bars = [("Mitosis", 14 / 18, GREEN, "14/18"), ("Plain RAG", 11 / 18, MUTED, "11/18")]
for j, (lab, frac, col, raw) in enumerate(bars):
    by = 4.55 + j * 0.7
    text(s, 1.1, by, 1.3, 0.35, lab, size=13, color=TEXT, anchor=MSO_ANCHOR.MIDDLE)
    shape(s, MSO_SHAPE.RECTANGLE, 2.45, by + 0.06, 3.0, 0.24, PANEL2, None)
    shape(s, MSO_SHAPE.RECTANGLE, 2.45, by + 0.06, 3.0 * frac, 0.24, col, None)
    text(s, 5.5, by, 0.7, 0.35, raw, size=12, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)

tiles = [
    ("19/21", "contradictions caught", "planted in the corpus; plain RAG: 0", CYAN),
    ("0", "client data leaks", "access control before retrieval", GREEN),
    ("0 vs 1", "prompt-injection leaks", "Mitosis vs plain RAG", AMBER),
    ("85%", "of routing without an LLM call", "rules + local embeddings", VIOLET),
]
tw, th, gx, gy = 2.95, 1.9, 0.3, 0.3
for k, (big, lab, sub, col) in enumerate(tiles):
    tx = 6.7 + (k % 2) * (tw + gx)
    ty = 2.1 + (k // 2) * (th + gy)
    card(s, tx, ty, tw, th)
    text(s, tx + 0.3, ty + 0.2, tw - 0.5, 0.8, big, size=40, bold=True, color=col)
    text(s, tx + 0.3, ty + 1.0, tw - 0.5, 0.4, lab, size=13, color=TEXT)
    text(s, tx + 0.3, ty + 1.38, tw - 0.5, 0.35, sub, size=11, color=MUTED)
text(s, 0.7, 6.55, 12, 0.4,
     [[("Security  ", {"color": CYAN, "bold": True}),
       ("Signed tokens, role checks on every route, per-client views (no IDOR), injections quarantined, "
        "PII redacted at ingest.", {"color": MUTED})]], size=14)
s.notes_slide.notes_text_frame.text = (
    "We tested it on 102 documents with 21 planted contradictions. Mitosis answered 78% of the golden "
    "questions correctly, against 61% for a plain chatbot on the same documents. It surfaced 19 of the 21 "
    "contradictions before anyone asked. Client data stays with the client: zero leaks. And a document that "
    "tries to give the AI instructions gets quarantined; the plain chatbot followed it once. 85% of routing "
    "needs no AI call at all.")

# ---------------------------------------------------------------- slide 5
s = new_slide()
kicker(s, 5, "Why SD Worx, why now")
text(s, 0.7, 0.95, 12, 0.9, "The trust layer under SD Worx’s agents.", size=40, bold=True)

# left: layer stack
layers = [
    ("SD Worx AI agents", "agentic payroll, since June 2026", MUTED, PANEL),
    ("Mitosis", "which answer to trust, and who owns it", CYAN, "0B2A2E"),
    ("Documents, Teams, tickets, client configs", "the knowledge as it really is", MUTED, PANEL),
]
lx, lw_, ly = 0.7, 5.8, 2.15
for i, (name, sub, col, fill) in enumerate(layers):
    y = ly + i * 1.3
    card(s, lx, y, lw_, 1.05, fill=fill, line=CYAN if i == 1 else LINE)
    text(s, lx + 0.35, y + 0.18, lw_ - 0.7, 0.4, name, size=20 if i == 1 else 16, bold=True,
         color=col if i == 1 else TEXT)
    text(s, lx + 0.35, y + 0.6, lw_ - 0.7, 0.35, sub, size=12, color=MUTED)

# right: rollout path as dividing cells
rx = 7.3
text(s, rx, 2.15, 5.3, 0.35, "ROLLOUT", size=12, color=MUTED, bold=True)
path = [
    ("PC 200 team", "one joint committee, one team"),
    ("All joint committees", "Belgium"),
    ("27 countries", "every SD Worx payroll desk"),
    ("mysdworx", "client HR admins, same verified answers"),
]
for i, (name, sub) in enumerate(path):
    y = 2.7 + i * 0.85
    rr = 0.14 + i * 0.03
    if i < len(path) - 1:
        connector(s, rx + 0.3, y + 0.25, rx + 0.3, y + 0.85 + 0.25 - 0.02, LINE, 1.5)
    cell(s, rx + 0.3, y + 0.25, rr, GREEN if i == 0 else CYAN, lw=2)
    text(s, rx + 0.8, y + 0.02, 4.6, 0.35, name, size=16, bold=True)
    text(s, rx + 0.8, y + 0.38, 4.6, 0.3, sub, size=12, color=MUTED)

text(s, 0.7, 6.3, 12, 0.7,
     [[("Find it. ", {"color": TEXT}), ("Understand it. ", {"color": TEXT}), ("Trust it.", {"color": CYAN})]],
     size=34, bold=True)
s.notes_slide.notes_text_frame.text = (
    "In June 2026 SD Worx put its payroll consultants on AI agents that retrieve knowledge. More AI answers "
    "means more answers that sound sure and are wrong. We start with one joint committee and one team, PC 200, "
    "then all joint committees, then the 27 countries, then client-facing inside mysdworx. "
    "SD Worx already has agents that find information. Mitosis is what lets people trust what they find. "
    "Find it. Understand it. Trust it.")

prs.save(OUT)
print(OUT)
